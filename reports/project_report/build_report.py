"""Build the project report (DOCX and PDF) from its Markdown source.

    python reports/project_report/build_report.py

Input:  reports/project_report/CS3011_Group52_Project_Report.md (the editable source)
Output: reports/project_report/CS3011_Group52_Project_Report.docx
        reports/project_report/CS3011_Group52_Project_Report.pdf

Needs pandoc and python-docx (pip install python-docx). The PDF step and the
filled-in table of contents need LibreOffice with its Python bridge (`import uno`,
available on Linux with the libreoffice-script-provider-python package). Without
it, the DOCX is still written and Word fills in the contents when you right-click
it and choose Update Field.

This script only lays out text that is already in the Markdown file. It does not
read or change any result; every number comes from the source file, which must
agree with reports/RESULTS.md.
"""

from __future__ import annotations

import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "CS3011_Group52_Project_Report.md"
OUT_DOCX = HERE / "CS3011_Group52_Project_Report.docx"
OUT_PDF = HERE / "CS3011_Group52_Project_Report.pdf"

RUNNING_HEAD = "CS3011 Group 52 · AI Based Public Transport Demand and Occupancy Prediction"
BODY_FONT, HEAD_FONT, CODE_FONT = "Calibri", "Cambria", "Consolas"
DARK, MID, GREY = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x2E, 0x4A, 0x86), RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "DCE3F0"
# Headings that start a new page; other sections follow on, so short endings do not leave empty pages.
NEW_PAGE = ("Abstract", "1. ", "References", "Appendix A")


# --------------------------------------------------------------------------- source

def split_source(text: str) -> tuple[list[str], str]:
    """Return (title-page lines, body markdown without front matter or title page)."""
    text = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)           # YAML front matter
    text = re.sub(r"<!--(?!\s*titlepage).*?-->", "", text, flags=re.S)  # build comments
    m = re.search(r"<!-- titlepage:start.*?-->(.*?)<!-- titlepage:end -->", text, flags=re.S)
    if not m:
        raise SystemExit("title page markers not found in the source")
    lines = [ln.strip() for ln in m.group(1).strip().splitlines() if ln.strip()]
    return lines, text[m.end():]


def plain(s: str) -> str:
    s = re.sub(r"\*\*(.*?)\*\*", r"\1", s)
    s = re.sub(r"<(https?://[^>]+)>", r"\1", s)
    return s.replace("`", "").lstrip("# ").strip()


# --------------------------------------------------------------------------- styles

def set_font(style, name: str, size: float | None = None, bold=None, italic=None, color=None) -> None:
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(a), None)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def make_reference(path: Path) -> None:
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout
    path.write_bytes(data)
    doc = Document(str(path))
    # Look styles up by their stored name: python-docx maps "Heading 1" to the
    # lowercase built-in name, which pandoc's reference file does not use.
    st = {s.name: s for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in st:
            s = st[name]
            set_font(s, BODY_FONT, 11 if name != "Compact" else 10.5, color=RGBColor(0, 0, 0))
            pf = s.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(6 if name != "Compact" else 2)
            pf.line_spacing = 1.15
    for name, size, color, before, after in (("Heading 1", 17, DARK, 0, 12), ("Heading 2", 13, MID, 14, 6),
                                             ("Heading 3", 11.5, MID, 10, 4)):
        s = st[name]
        set_font(s, HEAD_FONT, size, bold=True, italic=False, color=color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            s.paragraph_format.space_before = Pt(22)
    for name in ("Image Caption", "Table Caption", "Caption"):
        if name in st:
            set_font(st[name], BODY_FONT, 9.5, italic=True, color=GREY)
            st[name].paragraph_format.space_after = Pt(12)
    if "Source Code" in st:
        set_font(st["Source Code"], CODE_FONT, 9)
    if "Verbatim Char" in st:
        set_font(st["Verbatim Char"], CODE_FONT, 9.5)
    if "Hyperlink" in st:
        st["Hyperlink"].font.color.rgb = MID
    for name in ("Title", "Subtitle", "Date"):
        if name in st:
            set_font(st[name], HEAD_FONT)
    doc.save(str(path))


# --------------------------------------------------------------------------- layout helpers

def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table_borders(table) -> None:
    tblpr = table._tbl.tblPr
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, size in (("top", 8), ("left", 0), ("bottom", 8), ("right", 0), ("insideH", 4), ("insideV", 0)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if size else "nil")
        if size:
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), "8496B0")
        borders.append(el)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 30), ("left", 80), ("bottom", 30), ("right", 80)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


TEXT_WIDTH_CM = 21.0 - 2 * 2.4


def column_widths(t) -> list[float]:
    """Column widths (cm) in proportion to content: long text columns get more room."""
    # Every column is at least as wide as its longest word, so no word is split;
    # the remaining width goes to the columns with the longest text.
    char_cm, pad_cm = 0.2, 0.35
    mins, wants = [], []
    for j in range(len(t.columns)):
        texts = [row.cells[j].text for row in t.rows if j < len(row.cells)]
        longest_word = max((len(w) for tx in texts for w in tx.split()), default=3)
        longest_text = max((len(tx) for tx in texts[1:]), default=longest_word)
        mins.append(longest_word * char_cm + pad_cm)
        wants.append(max(mins[-1], min(longest_text, 70) * char_cm + pad_cm))
    if sum(wants) <= TEXT_WIDTH_CM:
        return [w * TEXT_WIDTH_CM / sum(wants) for w in wants]
    if sum(mins) >= TEXT_WIDTH_CM:
        return [m * TEXT_WIDTH_CM / sum(mins) for m in mins]
    extra = TEXT_WIDTH_CM - sum(mins)
    gaps = [w - m for w, m in zip(wants, mins)]
    return [m + extra * g / sum(gaps) for m, g in zip(mins, gaps)]


def set_widths(t, widths: list[float]) -> None:
    tblpr = t._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(sum(widths) / 2.54 * 1440)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)


def format_tables(doc) -> None:
    from docx.text.paragraph import Paragraph
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t)
        set_widths(t, column_widths(t))
        # keep the bold "Table N." line with its table, and leave space after the table
        prev = t._tbl.getprevious()
        if prev is not None and prev.tag == qn("w:p"):
            Paragraph(prev, t._parent).paragraph_format.keep_with_next = True
        nxt = t._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            Paragraph(nxt, t._parent).paragraph_format.space_before = Pt(8)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trpr.append(cant)
            if i == 0:
                hdr = OxmlElement("w:tblHeader")
                trpr.append(hdr)
            for cell in row.cells:
                if i == 0:
                    shade(cell, HEADER_FILL)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if i == 0:
                            r.font.bold = True


def add_field(paragraph, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        r = OxmlElement("w:r")
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r.append(it)
        paragraph._p.append(r)


def header_footer(doc) -> None:
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = RUNNING_HEAD
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size, r.font.color.rgb, r.font.name = Pt(8.5), GREY, BODY_FONT
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Page ")
    run.font.size, run.font.color.rgb = Pt(9), GREY
    add_field(fp, "PAGE")
    for r in fp.runs:
        r.font.size = Pt(9)


def centred(doc_body, before_el, text: str, size: float, bold=False, color=None, font=BODY_FONT,
            space_after: float = 6, italic=False):
    """Insert a centred paragraph before `before_el` and return it."""
    p = before_el.insert_paragraph_before("")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
    if color is not None:
        r.font.color.rgb = color
    return p


def title_page_and_contents(doc, lines: list[str]) -> None:
    first = doc.paragraphs[0]
    table_rows = [ln for ln in lines if ln.startswith("|") and not set(ln) <= set("|- ")]
    text_lines = [ln for ln in lines if not ln.startswith("|")]
    org, course, title, group, *rest = text_lines

    spacer = centred(doc, first, "", 10, space_after=40)
    centred(doc, first, plain(org), 12, bold=True, color=DARK, space_after=4)
    centred(doc, first, plain(course), 11, color=GREY, space_after=60)
    centred(doc, first, plain(title), 26, bold=True, color=DARK, font=HEAD_FONT, space_after=18)
    centred(doc, first, plain(group), 13, bold=True, color=MID, space_after=28)

    # team table
    rows = [[plain(c) for c in r.strip("|").split("|")] for r in table_rows]
    p_anchor = centred(doc, first, "", 4, space_after=0)
    tbl = doc.add_table(rows=len(rows), cols=2)
    p_anchor._p.addnext(tbl._tbl)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = val
            for r in cell.paragraphs[0].runs:
                r.font.size, r.font.name, r.font.bold = Pt(10.5), BODY_FONT, i == 0
    tbl.columns[0].width, tbl.columns[1].width = Cm(6.5), Cm(7.5)
    for row in tbl.rows:
        row.cells[0].width, row.cells[1].width = Cm(6.5), Cm(7.5)

    after_tbl = centred(doc, first, "", 10, space_after=50)
    for ln in rest:
        centred(doc, first, plain(ln), 10, color=GREY, space_after=4)

    # contents page
    br = first.insert_paragraph_before("")
    br.add_run().add_break(WD_BREAK.PAGE)
    h = first.insert_paragraph_before("")
    hr = h.add_run("Contents")
    hr.font.size, hr.font.bold, hr.font.name, hr.font.color.rgb = Pt(17), True, HEAD_FONT, DARK
    h.paragraph_format.space_after = Pt(12)
    toc = first.insert_paragraph_before("")
    add_field(toc, r'TOC \o "1-2" \h \z \u')
    del spacer, after_tbl
    # no forced page break needed: Heading 1 starts a new page


def new_pages(doc, starts: tuple[str, ...]) -> None:
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text.strip().startswith(starts):
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(0)


# --------------------------------------------------------------------------- LibreOffice

def libreoffice_update_and_pdf(docx: Path, pdf: Path) -> list[tuple[str, str]] | None:
    """Update the table of contents and export the PDF. Returns the TOC entries."""
    try:
        import uno  # noqa: F401
        from com.sun.star.beans import PropertyValue  # noqa: F401
    except ImportError:
        return None
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with socket.socket() as s:
        s.bind(("localhost", 0))
        port = s.getsockname()[1]
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    proc = subprocess.Popen([soffice, "--headless", "--invisible", "--norestore", "--nologo",
                             f"-env:UserInstallation=file://{profile}",
                             f"--accept=socket,host=localhost,port={port};urp;"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        import uno
        from com.sun.star.beans import PropertyValue

        def prop(name, value):
            p = PropertyValue()
            p.Name, p.Value = name, value
            return p

        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(60):
            try:
                ctx = resolver.resolve(f"uno:socket,host=localhost,port={port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        else:
            raise RuntimeError("could not connect to LibreOffice")
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(str(docx)), "_blank", 0,
                                           (prop("Hidden", True),))
        idx = doc.getDocumentIndexes()
        if idx.getCount() == 0:
            # LibreOffice does not turn Word's TOC field into an index on import, so
            # insert its own table of contents right after the "Contents" heading.
            text = doc.getText()
            paras = text.createEnumeration()
            while paras.hasMoreElements():
                par = paras.nextElement()
                if par.supportsService("com.sun.star.text.Paragraph") and par.getString().strip() == "Contents":
                    toc = doc.createInstance("com.sun.star.text.ContentIndex")
                    toc.CreateFromOutline = True
                    toc.Level = 2
                    toc.Title = ""
                    text.insertTextContent(text.createTextCursorByRange(par.getEnd()), toc, False)
                    break
            idx = doc.getDocumentIndexes()
        entries: list[tuple[str, str]] = []
        for _ in range(2):            # second pass settles page numbers if the TOC grows
            for i in range(idx.getCount()):
                idx.getByIndex(i).update()
            doc.refresh()
        if idx.getCount():
            text = idx.getByIndex(0).getAnchor().getString()
            for line in text.splitlines():
                if "\t" in line:
                    t, pg = line.rsplit("\t", 1)
                    entries.append((t.strip(), pg.strip()))
        doc.storeToURL(uno.systemPathToFileUrl(str(pdf)), (prop("FilterName", "writer_pdf_Export"),))
        doc.close(True)
        return entries
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=20)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(profile, ignore_errors=True)


def cache_toc(docx: Path, entries: list[tuple[str, str]]) -> None:
    """Write the computed TOC lines into the field's result, so viewers that do not
    update fields (LibreOffice, Google Docs, previews) still show the contents."""
    doc = Document(str(docx))
    for p in doc.paragraphs:
        if any(it.text and it.text.strip().startswith("TOC") for it in p._p.iter(qn("w:instrText"))):
            toc_p = p
            break
    else:
        return
    runs = list(toc_p._p.iter(qn("w:r")))
    end_run = runs[-1]                      # fldChar end
    toc_p._p.remove(end_run)
    usable = Cm(21.0 - 2 * 2.4)

    def fill(par, title, page, level):
        par.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        par.paragraph_format.left_indent = Cm(0 if level == 1 else 0.6)
        par.paragraph_format.space_after = Pt(3 if level == 1 else 1)
        par.paragraph_format.space_before = Pt(5 if level == 1 else 0)
        r = par.add_run(f"{title}\t{page}")
        r.font.size, r.font.bold, r.font.name = Pt(10.5 if level == 1 else 10), level == 1, BODY_FONT

    headings = {pp.text.strip(): (1 if pp.style.name == "Heading 1" else 2)
                for pp in doc.paragraphs if pp.style.name in ("Heading 1", "Heading 2")}
    prev = toc_p
    for k, (title, page) in enumerate(entries):
        level = headings.get(title, 2)
        if k == 0:
            fill(toc_p, title, page, level)
            continue
        newp = OxmlElement("w:p")
        prev._p.addnext(newp)
        from docx.text.paragraph import Paragraph
        prev = Paragraph(newp, toc_p._parent)
        fill(prev, title, page, level)
    prev._p.append(end_run)
    doc.save(str(docx))


# --------------------------------------------------------------------------- schema order

# Child order required by the WordprocessingML schema (ECMA-376). pandoc and the
# edits above can leave children out of order; Word is strict about some of these.
ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd "
           "tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc "
           "textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof "
           "snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText "
           "vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc "
             "tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "style": "name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed qFormat "
             "locked personal personalCompose personalReply rsid pPr rPr tblPr trPr tcPr tblStylePr",
    "settings": "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
                "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData "
                "embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins "
                "alignBordersAndEdges bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop "
                "hideSpellingErrors hideGrammaticalErrors activeWritingStyle proofState formsDesign "
                "attachedTemplate linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge "
                "revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection "
                "autoFormatOverride styleLockTheme styleLockQFSet defaultTabStop autoHyphenation "
                "consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope summaryLength "
                "clickAndTypeStyle defaultTableStyle evenAndOddHeaders bookFoldRevPrinting bookFoldPrinting "
                "bookFoldPrintingSheets drawingGridHorizontalSpacing drawingGridVerticalSpacing "
                "displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery "
                "doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin drawingGridVerticalOrigin "
                "doNotShadeFormData noPunctuationKerning characterSpacingControl printTwoOnOne "
                "strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore savePreviewPicture "
                "doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText "
                "doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags "
                "alwaysMergeEmptyNamespace updateFields hdrShapeDefaults footnotePr endnotePr compat docVars "
                "rsids mathPr attachedSchema themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats "
                "doNotAutoCompressPictures forceUpgrade captions readModeInkLockDown smartTagType "
                "schemaLibrary shapeDefaults doNotEmbedSmartTags decimalSymbol listSeparator",
}
ORDER = {k: {name: i for i, name in enumerate(v.split())} for k, v in ORDER.items()}


def normalise_xml(docx: Path) -> None:
    import zipfile
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    parts = ("word/document.xml", "word/styles.xml", "word/settings.xml", "word/numbering.xml")
    tmp = docx.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in parts:
                root = etree.fromstring(data)
                for parent_name, rank in ORDER.items():
                    for el in root.iter(W + parent_name):
                        kids = list(el)
                        last = len(rank)
                        # rank by local name, whatever the namespace (settings holds m:mathPr)
                        known = sorted((k for k in kids if isinstance(k.tag, str)),
                                       key=lambda k: rank.get(etree.QName(k).localname, last))
                        others = [k for k in kids if not isinstance(k.tag, str)]
                        for k in kids:
                            el.remove(k)
                        for k in known + others:
                            if k.tail and k.tail.strip():   # stray text between elements
                                k.tail = None
                            el.append(k)
                for el in root.iter(W + "nsid"):          # hex IDs must have 8 digits
                    v = el.get(W + "val")
                    if v and len(v) < 8:
                        el.set(W + "val", v.rjust(8, "0"))
                for el in root.iter(W + "pgMar"):         # header, footer, gutter are required
                    for a, default in (("header", "709"), ("footer", "709"), ("gutter", "0")):
                        if el.get(W + a) is None:
                            el.set(W + a, default)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    tmp.replace(docx)


# --------------------------------------------------------------------------- main

def main() -> int:
    lines, body = split_source(SRC.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        ref = tmp / "reference.docx"
        make_reference(ref)
        md = tmp / "body.md"
        md.write_text(body, encoding="utf-8")
        raw = tmp / "raw.docx"
        subprocess.run(["pandoc", str(md), "-f", "markdown", "-t", "docx", "--reference-doc", str(ref),
                        "--resource-path", str(HERE), "-o", str(raw)], check=True, cwd=HERE)
        doc = Document(str(raw))
        title_page_and_contents(doc, lines)
        new_pages(doc, NEW_PAGE)
        format_tables(doc)
        header_footer(doc)
        doc.core_properties.title = "AI Based Public Transport Demand and Occupancy Prediction"
        doc.core_properties.subject = "CS3011 Artificial Intelligence, Group 52, project report"
        doc.core_properties.author = "CS3011 Group 52"
        doc.save(str(OUT_DOCX))
    normalise_xml(OUT_DOCX)

    entries = libreoffice_update_and_pdf(OUT_DOCX, OUT_PDF)
    if entries is None:
        print(f"Wrote {OUT_DOCX.name}. LibreOffice's Python bridge is not available, so no PDF was made and "
              "the contents are filled in by Word (right-click the contents, Update Field).")
        return 0
    cache_toc(OUT_DOCX, entries)
    normalise_xml(OUT_DOCX)
    print(f"Wrote {OUT_DOCX.name} and {OUT_PDF.name} ({len(entries)} contents entries).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
